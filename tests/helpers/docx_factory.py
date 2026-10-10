from io import BytesIO
from posixpath import relpath
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def paragraph_table_docx(before: str, rows: tuple[tuple[str, ...], ...], after: str) -> bytes:
    document = Document()
    _ = document.add_paragraph(before)
    if rows:
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for row, values in zip(table.rows, rows, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    _ = document.add_paragraph(after)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def rewrite_docx(data: bytes, replacements: dict[str, bytes]) -> bytes:
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            target.writestr(entry.filename, replacements.get(entry.filename, source.read(entry)))
        for name, value in replacements.items():
            if name not in source.namelist():
                entry = ZipInfo(name)
                entry.filename = name
                entry.orig_filename = name
                target.writestr(entry, value)
    return output.getvalue()


def body_xml(content: str) -> bytes:
    return (
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f"<w:body>{content}<w:sectPr/></w:body></w:document>"
    ).encode()


def with_header(data: bytes, part: str, content: bytes) -> bytes:
    with ZipFile(BytesIO(data)) as source:
        types = source.read("[Content_Types].xml")
        relationships = source.read("word/_rels/document.xml.rels")
        document = source.read("word/document.xml")
    namespace = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/header"
    media = "application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"
    relationship = (
        f'<Relationship Id="story" Type="{namespace}" Target="{relpath(part, "word")}"/>'.encode()
    )
    override = f'<Override PartName="/{part}" ContentType="{media}"/>'.encode()
    offset = document.index(b">", document.index(b"<w:sectPr")) + 1
    document = (
        document[:offset]
        + b'<w:headerReference w:type="default" r:id="story"/>'
        + document[offset:]
    )
    return rewrite_docx(
        data,
        {
            part: content,
            "word/document.xml": document,
            "[Content_Types].xml": types.replace(b"</Types>", override + b"</Types>"),
            "word/_rels/document.xml.rels": relationships.replace(
                b"</Relationships>", relationship + b"</Relationships>"
            ),
        },
    )
